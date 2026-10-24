from docx import Document
from zipfile import ZipFile
from lxml import etree
p='output/DBTHON_2026_Assessment_7_Project_Report.docx'
d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables))
print('Jev table row',any('Jev decision support' in cell.text for t in d.tables for row in t.rows for cell in row.cells))
print('headings',[x.text for x in d.paragraphs if x.style.name.startswith('Heading')])
with ZipFile(p) as z:
    root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    colors=[e.get('{%s}val'%ns['w']) for e in root.xpath('.//w:color',namespaces=ns) if e.get('{%s}val'%ns['w']) not in ('auto',None)]
    fills=[e.get('{%s}fill'%ns['w']) for e in root.xpath('.//w:shd',namespaces=ns) if e.get('{%s}fill'%ns['w']) not in ('auto',None)]
    def gray(c): return len(c)==6 and c[0:2]==c[2:4]==c[4:6]
    print('explicit text colors',sorted(set(colors)))
    print('non-gray text colors',[c for c in colors if not gray(c)])
    print('fills',sorted(set(fills)))
